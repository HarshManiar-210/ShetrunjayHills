#!/usr/bin/env python3
"""Turn the client's per-layer statistics spreadsheets into overlay_class_stats rows.

The client delivered one .xlsx per theme (the file name says which), holding
the official area and share of each class, per year where the theme has years:

    Forest Cover.xlsx        forest_cover_<year>, one column pair per year
    LULC.xlsx                lulc_<year>, likewise
    LULC Drone.xlsx          lulc_2026 (the drone-classified current LULC)
    Fragmentation.xlsx       fragmentation_<year>
    Vegetation Change.xlsx   vegetation_change_<from>_<to>, one sheet each
    Tree Density Drone.xlsx  treeDensity
    Forest Cover FSI.xlsx    forestCoverFSI (vector)
    Forest Tpe FSI.xlsx      forestTypeFSI (vector)

Each class is written with the value its theme's legend already uses (the
raster legends in apps/web/lib/legend-config.ts, the FSI layers' categories
in init.sql), because that value is how the frontend finds the class's colour
and display name. A class name this script doesn't recognise stops it rather
than guessing -- a mismatched key would show as an unclassified grey row.

Areas are hectares and percentages 0-100, both exactly as delivered. Rows with
no area (e.g. Fragmentation's Core 250-500 acres, only mapped in 1980) are
dropped rather than stored as zero.

Usage:
    pip install openpyxl
    python tools/prepare-class-stats.py <folder of .xlsx>  > rows.sql

The output is the VALUES list used by both infra/postgis-init/init.sql and
infra/migrations/015-overlay-class-stats.sql; paste it into each.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

import openpyxl

# --- class name -> legend value ----------------------------------------------

DENSITY = {
    "very dense forest": "1",
    "moderately dense forest": "2",
    "open forest": "3",
    "scrub": "4",
    "non-forest": "5",
}
DENSITY_LABEL = {
    "1": "Very Dense Forest",
    "2": "Moderately Dense Forest",
    "3": "Open Forest",
    "4": "Scrub",
    "5": "Non Forest",
}

LULC = {
    "barren": "1",
    "builtup": "2",
    "dense vegetation": "3",
    "dense": "3",
    "scrub/sparse vegetation": "4",
    "sparse": "4",
    "water": "5",
    "waterbody": "5",
}
LULC_LABEL = {
    "1": "Barren",
    "2": "Builtup",
    "3": "Dense Vegetation",
    "4": "Scrub / Sparse Vegetation",
    "5": "Waterbody",
}

FRAGMENTATION = {
    "patch": ("patch", "Patch"),
    "edge": ("edge", "Edge"),
    "perforated": ("perforated", "Perforated"),
    "core(<250 acres)": ("core-small", "Core (<250 acres)"),
    "core(250-500 acres)": ("core-medium", "Core (250–500 acres)"),
    "core(>500 acres)": ("core-large", "Core (>500 acres)"),
}

DENSITY_CODES = {"VDF", "MDF", "OF", "SCRUB", "NF"}
# Densest first. A transition's group follows from its two codes: towards a
# denser class is Improvement, towards a sparser one Degradation. Derived
# rather than read from the sheet, whose headings misfile some rows (e.g.
# NF-SCRUB under Degradation in 1980-1989).
DENSITY_RANK = {"VDF": 0, "MDF": 1, "OF": 2, "SCRUB": 3, "NF": 4}
GROUP_ORDER = ("Improvement", "Degradation", "Stable")


def transition_group(frm: str, to: str) -> str:
    d = DENSITY_RANK[to] - DENSITY_RANK[frm]
    return "Improvement" if d < 0 else "Degradation" if d > 0 else "Stable"
# One typo in the delivery ("VD-SCRUB", 2018 over 2008).
CODE_FIXES = {"VD": "VDF"}

# Delivered FSI names are the features' own Type values, which is what the
# layers' categories are keyed on -- so they pass through unchanged.
FSI_COVER_LABEL = {
    "MODERATELY DENSE FOREST (Tree Canopy density 40% & above but < 70%)": "Moderately Dense Forest",
    "OPEN FOREST (Tree Canopy density 10% & above but < 40%)": "Open Forest",
    "SCRUB (Tree Canopy density < 10%)": "Scrub",
    "WATER": "Water",
}
FSI_TYPE_LABEL = {
    "3B/C2 Southern moist mixed deciduous forest": "3B/C2 Southern moist mixed deciduous forest",
    "5/DS4 Dry Grassland": "5/DS4 Dry Grassland",
    "5/E1 Anogeissus pendula Forest": "5/E1 Anogeissus pendula Forest",
    "5/E 8c Salvadora-T amarix scrub": "5/E 8c Salvadora-Tamarix scrub",
    "Acacia senegal forest": "Acacia senegal forest",
    "6/E4 Salvadora scrub": "6/E4 Salvadora scrub",
    "Water": "Water",
}

# (overlay key, value, label, group, area ha, percentage, sort order)
Row = tuple[str, str, str, "str | None", float, float, int]


def rows_of(path: Path, sheet=None) -> list[tuple]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    return [r for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]


def lookup(table: dict, name: str, source: str):
    key = re.sub(r"\s+", " ", str(name)).strip().lower()
    if key not in table:
        sys.exit(f"{source}: unrecognised class {name!r}")
    return table[key]


def yearwise(path: Path, prefix: str, name_col: int, table: dict, labels: dict | None) -> list[Row]:
    """A sheet with one Area/Percentage column pair per year."""
    rows = rows_of(path)
    header = rows[0]
    out: list[Row] = []
    for col, year in enumerate(header):
        if not isinstance(year, (int, float)) or year < 1900:
            continue
        order = 0
        for r in rows[2:]:
            name = r[name_col]
            if name is None or str(name).strip().lower() == "total":
                continue
            area, pct = r[col], r[col + 1]
            if area is None:
                continue
            hit = lookup(table, name, path.name)
            value, label = (hit, labels[hit]) if labels else hit
            order += 1
            out.append((f"{prefix}_{int(year)}", value, label, None, area, pct, order))
    return out


def single(path: Path, key: str, table: dict, labels: dict) -> list[Row]:
    """Class / area only, no percentage column: shares are taken of the classes' sum."""
    classes = [
        (r[0], r[1]) for r in rows_of(path)[1:]
        if r[0] is not None and str(r[0]).strip().lower() != "total"
    ]
    total = sum(area for _, area in classes)
    out: list[Row] = []
    for order, (name, area) in enumerate(classes, start=1):
        value = lookup(table, name, path.name)
        out.append((key, value, labels[value], None, area, area / total * 100, order))
    return out


def fsi(path: Path, key: str, labels: dict) -> list[Row]:
    out: list[Row] = []
    for r in rows_of(path)[1:]:
        name = r[1]
        if name is None:
            continue
        name = str(name).strip()
        if name not in labels:
            sys.exit(f"{path.name}: unrecognised class {name!r}")
        out.append((key, name, labels[name], None, r[2], r[3], len(out) + 1))
    return out


def vegetation_change(path: Path) -> list[Row]:
    out: list[Row] = []
    for sheet in openpyxl.load_workbook(path, read_only=True).sheetnames:
        m = re.fullmatch(r"(\d{4}) over (\d{4})", sheet.strip())
        if not m:
            sys.exit(f"{path.name}: unexpected sheet {sheet!r}")
        key = f"vegetation_change_{m[2]}_{m[1]}"
        sheet_rows = []
        for r in rows_of(path, sheet)[1:]:
            if r[1] is None or r[2] is None:
                continue
            codes = [CODE_FIXES.get(c, c) for c in re.split(r"\s*-\s*", str(r[1]).strip().upper())]
            if len(codes) != 2 or not set(codes) <= DENSITY_CODES:
                sys.exit(f"{path.name}/{sheet}: unrecognised transition {r[1]!r}")
            value = "-".join(codes)
            sheet_rows.append((value, transition_group(*codes), r[2], r[3]))
        # The panel heads each run of one group, so a group's rows must be
        # contiguous; the sort is stable, keeping the delivered order within one.
        sheet_rows.sort(key=lambda row: GROUP_ORDER.index(row[1]))
        for order, (value, group, area, pct) in enumerate(sheet_rows, 1):
            out.append((key, value, value.replace("-", " → "), group, area, pct, order))
    return out


def sql_str(s: "str | None") -> str:
    return "NULL" if s is None else "'" + s.replace("'", "''") + "'"


def main() -> int:
    if len(sys.argv) != 2:
        print(__doc__, file=sys.stderr)
        return 2
    d = Path(sys.argv[1])

    rows: list[Row] = []
    rows += yearwise(d / "Forest Cover.xlsx", "forest_cover", 0, DENSITY, DENSITY_LABEL)
    rows += yearwise(d / "LULC.xlsx", "lulc", 0, LULC, LULC_LABEL)
    rows += single(d / "LULC Drone.xlsx", "lulc_2026", LULC, LULC_LABEL)
    rows += yearwise(d / "Fragmentation.xlsx", "fragmentation", 1, FRAGMENTATION, None)
    rows += vegetation_change(d / "Vegetation Change.xlsx")
    rows += single(d / "Tree Density Drone.xlsx", "treeDensity", DENSITY, DENSITY_LABEL)
    rows += fsi(d / "Forest Cover FSI.xlsx", "forestCoverFSI", FSI_COVER_LABEL)
    rows += fsi(d / "Forest Tpe FSI.xlsx", "forestTypeFSI", FSI_TYPE_LABEL)

    lines = [
        f"    ({sql_str(k)}, {sql_str(v)}, {sql_str(l)}, {sql_str(g)}, {round(a, 4)}, {round(p, 4)}, {o})"
        for k, v, l, g, a, p, o in rows
    ]
    print(",\n".join(lines))
    return 0


if __name__ == "__main__":
    sys.exit(main())
